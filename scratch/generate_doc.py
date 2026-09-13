import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_styled_document(output_path):
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x22, 0x32, 0x4A) # text-strong

    # Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("ComputeBestSpecs")
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x12, 0x6B, 0xD6) # brand-primary
    title_p.paragraph_format.space_after = Pt(2)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("UI Design System, Component Architecture & Multi-App Concurrency Specification")
    sub_run.font.size = Pt(14)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(0x22, 0x32, 0x4A)
    sub_p.paragraph_format.space_after = Pt(6)

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run("Updated Edition (v2.0) • Complete Design Tokens, Light/System/Dark Modes, Application Catalog & Concurrency Architecture")
    meta_run.font.size = Pt(9.5)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(0x7D, 0x8C, 0xA3)
    meta_p.paragraph_format.space_after = Pt(14)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x12, 0x6B, 0xD6)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x22, 0x32, 0x4A)
        return p

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = RGBColor(0x22, 0x32, 0x4A)
        r_text = p.add_run(text)
        r_text.font.color.rgb = RGBColor(0x53, 0x67, 0x83)
        return p

    def add_bullet(bold_txt, body_txt):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r_b = p.add_run(bold_txt + ": ")
        r_b.bold = True
        r_b.font.color.rgb = RGBColor(0x22, 0x32, 0x4A)
        r_t = p.add_run(body_txt)
        r_t.font.color.rgb = RGBColor(0x53, 0x67, 0x83)
        return p

    # --- SECTION 1 ---
    add_h1("1. Purpose and Non-Negotiable Design Rules")
    add_body("This specification is the single source of truth for the visual design system, interaction tokens, and component behavior across ComputeBestSpecs. It prevents accidental low-contrast combinations, ensures mathematical concurrency clarity, and establishes clean responsive patterns for mobile, tablet, and desktop viewports.")
    add_bullet("Default Experience", "Light mode is the default daylight experience (#FFFEFC ivory canvas). Dark mode (#0B0F17) is an equal OLED-optimized theme. System mode dynamically follows the OS.")
    add_bullet("Semantic Token Law", "All components consume CSS custom properties (e.g., var(--card-main), var(--text-strong)). Hard-coded hex colors in application components are strictly forbidden.")
    add_bullet("Selection vs. Semantic Status", "The vibrant orange token (#FF761A) is designated exclusively as an Active Selection Accent (e.g. chosen presets, active chips). It must never be used to represent warning or error states.")
    add_bullet("Dual-Mode Concurrency Clarity", "The multi-workload engine must clearly communicate whether calculations are running in Simultaneous Multitasking Mode (combined working sets + OS overhead) or Isolated Standalone Mode.")

    # --- SECTION 2 ---
    add_h1("2. Semantic Color Architecture & Design Tokens")
    add_body("The color architecture uses curated HSL/Hex tokens with distinct roles for Day and Night modes:")

    add_h2("2.1 Day / Light Mode Palette (Default)")
    
    # Table for Day Tokens
    table_day = doc.add_table(rows=1, cols=4)
    table_day.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_day.autofit = False

    hdr_cells = table_day.rows[0].cells
    hdr_titles = ["Token Name", "CSS Variable", "Hex Value", "Usage & Semantics"]
    col_widths = [Inches(1.5), Inches(1.8), Inches(1.0), Inches(2.5)]
    
    for idx, heading in enumerate(hdr_titles):
        hdr_cells[idx].text = heading
        hdr_cells[idx].width = col_widths[idx]
        set_cell_background(hdr_cells[idx], "126BD6")
        for p in hdr_cells[idx].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    day_data = [
        ("Canvas / Background", "--bg-main", "#FFFEFC", "Primary page background (crisp warm ivory)"),
        ("Secondary Well", "--bg-secondary", "#F6F3EE", "Well areas, search bars, and subtle backgrounds"),
        ("Card Surface", "--card-main", "#FFFEFC", "Main component cards with subtle border definition"),
        ("Card Elevated", "--card-elevated", "#F0F2F1", "Hover states, interactive chips, and dropdown menus"),
        ("Typography Strong", "--text-strong", "#22324A", "Primary headings, titles, and critical scores (deep navy)"),
        ("Typography Body", "--text-body", "#536783", "Paragraph body text, subheadings, and explanations"),
        ("Typography Muted", "--text-muted", "#7D8CA3", "Secondary labels, timestamps, and input placeholders"),
        ("Border Subtle", "--border-subtle", "#D7E1EA", "Standard card borders, dividers, and container edges"),
        ("Border Strong", "--border-strong", "#C3D1DF", "Active inputs, hover boundaries, and emphasized cards"),
        ("Brand Primary", "--brand-primary", "#126BD6", "Primary buttons, active indicators, and focus rings"),
        ("Brand Cyan", "--brand-cyan", "#1689A8", "RAM & memory capacity indicators"),
        ("Brand Violet", "--brand-violet", "#7046D9", "VRAM & graphics engine indicators"),
        ("Selection Accent", "--accent-selection", "#FF761A", "Highlighted choices, selected model chips, active radio"),
    ]

    for row_idx, (name, var, hex_code, desc) in enumerate(day_data):
        row_cells = table_day.add_row().cells
        bg_hex = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([name, var, hex_code, desc]):
            row_cells[i].text = val
            row_cells[i].width = col_widths[i]
            set_cell_background(row_cells[i], bg_hex)
            for p in row_cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(0x22, 0x32, 0x4A) if i < 3 else RGBColor(0x53, 0x67, 0x83)
                    if i == 1 or i == 2:
                        r.font.name = 'Courier New'

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("2.2 Night / Dark Mode Palette")
    add_body("The night palette provides true dark efficiency without eye-strain, maintaining full contrast alignment with day tokens:")
    add_bullet("Night Canvas", "#0B0F17 (Deep slate void)")
    add_bullet("Night Card Surface", "#111827 (Layer 1 elevated card)")
    add_bullet("Night Card Elevated", "#1F2937 (Layer 2 interactive hover / active wells)")
    add_bullet("Night Text Primary", "#F3F4F6 (Crisp high-contrast text)")
    add_bullet("Night Text Body", "#9CA3AF (Balanced neutral reading text)")
    add_bullet("Night Borders", "#1F2937 (Subtle) • #374151 (Strong interactive boundary)")

    # --- SECTION 3 ---
    add_h1("3. Explicit Theme Distinction: Light vs. System vs. Dark")
    add_body("ComputeBestSpecs provides three distinct appearance states with explicit visual cues across the entire application interface:")

    table_theme = doc.add_table(rows=1, cols=4)
    table_theme.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_theme.autofit = False

    t_hdr = table_theme.rows[0].cells
    t_cols = [Inches(1.2), Inches(1.3), Inches(1.5), Inches(2.8)]
    for idx, h in enumerate(["Theme Mode", "Nav Icon", "DOM Attributes", "Visual Behavior & Status"]):
        t_hdr[idx].text = h
        t_hdr[idx].width = t_cols[idx]
        set_cell_background(t_hdr[idx], "22324A")
        for p in t_hdr[idx].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    theme_data = [
        ("Light Mode", "Sun (Amber)", 'data-theme="light"\ndata-theme-preference="light"', "Forces fixed ivory daylight palette (#FFFEFC). Title: 'Theme: Light (Fixed)'."),
        ("Dark Mode", "Moon (Indigo)", 'data-theme="dark"\ndata-theme-preference="dark"', "Forces fixed deep slate OLED palette (#0B0F17). Title: 'Theme: Dark (Fixed)'."),
        ("System Mode", "Laptop + Live Dot", 'data-theme="[active]"\ndata-theme-preference="system"', "Dynamically syncs with OS color schedule. Features live sync indicator dot and subtitle 'Auto (OS is Dark/Light)'.")
    ]

    for row_idx, (mode, icon, attrs, behavior) in enumerate(theme_data):
        r_cells = table_theme.add_row().cells
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([mode, icon, attrs, behavior]):
            r_cells[i].text = val
            r_cells[i].width = t_cols[i]
            set_cell_background(r_cells[i], bg)
            for p in r_cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(0x22, 0x32, 0x4A)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SECTION 4 ---
    add_h1("4. Searchable Application Catalog & Concurrency Builder")
    add_body("The application selection architecture enables users to build realistic multi-app stacks reflecting authentic daily professional workflows:")
    add_bullet("Curated Application Catalog", "35+ pre-calibrated software profiles across 6 domains (Creative/Design, Dev/Eng, Local AI/Data, Gaming/3D, CAD/BIM, Daily Office).")
    add_bullet("Search & Filter Controls", "Instant real-time query filtering by application name or category, accompanied by category pill filters and memory demand badges (e.g. '~8.0GB RAM').")
    add_bullet("Workload Intensity Selector", "Per-app intensity adjustment (Light: casual editing/small repos; Medium: production workflows; Heavy: 8K rendering, huge compilations).")
    add_bullet("Instance Multipliers", "Granular quantity counters for virtualization & multi-instance tools (e.g. running 2× Android Emulators or multiple Docker containers simultaneously).")

    add_h2("4.1 Simultaneous Concurrency Tooltip & Popover Spec")
    add_body("Next to the 'Run Simultaneously' toggle in both WorkloadStackBuilder and SoftwarePicker, an interactive (i) button reveals an explanatory popover:")
    add_bullet("✓ Checked (Simultaneous Multitasking)", "Simulates concurrent execution when all selected applications run in parallel. Memory working sets and CPU threads are stacked with concurrency weights (1.0 Foreground, 0.65 Background) plus OS base reserve overhead.")
    add_bullet("◻ Unchecked (Single-App Standalone)", "Evaluates whether the hardware can execute each program individually in isolation, assuming background heavy software is closed prior to launching.")

    # --- SECTION 5 ---
    add_h1("5. Hardware Profile Selector & Extended Catalog")
    add_body("The hardware configuration system accommodates both fast 1-click preset discovery and exhaustive manual parameter customization:")
    add_bullet("40+ Curated Presets", "Covers Apple Silicon (M4, M4 Pro, M4 Max, M3, M2 Ultra), NVIDIA GeForce RTX (Blackwell RTX 5090 32GB, RTX 5080 16GB, RTX 4090, RTX 4080 Super), AMD Ryzen & Radeon, Intel Core Ultra & 14th Gen.")
    add_bullet("Extended Manual Memory", "Categorized RAM dropdown spans from 8GB up to 256GB with vendor classification (Apple Unified, Mainstream DDR4/DDR5, Workstation ECC).")
    add_bullet("Storage & Scratch Disks", "NVMe SSD options up to 8TB with dynamic free-space calculations.")
    add_bullet("Operating System Families", "16 categorized OS options across macOS (Sequoia 15, Sonoma 14, Ventura 13), Windows (11, 10), and Linux (Ubuntu 24.04/22.04, Fedora, Arch, Debian).")

    # --- SECTION 6 ---
    add_h1("6. Diagnostic Engine, Bottleneck Cards & Live Simulation")
    add_body("The evaluation system presents a clear 6-step diagnostic hierarchy:")
    add_bullet("Step 1: Match Score & Verdict", "Normalized 0-100 score with verdict tier ('Recommended', 'Playable with Constraints', 'Bottleneck Detected').")
    add_bullet("Step 2: Primary Bottleneck Card", "Isolates the specific limiting resource (e.g. RAM, VRAM, Storage) with exact installed vs peak required metrics.")
    add_bullet("Step 3: One-Click Upgrade Simulation", "'Simulate Upgrade' action button tests the impact of adding RAM or upgrading GPU in real time.")
    add_bullet("Step 4: System Resource Fingerprint", "Multi-axis visual radar map displaying capacity distribution across CPU, GPU, RAM, VRAM, and Storage.")
    add_bullet("Step 5: Interactive What-If Simulator", "Sliders and component swaps to model future hardware upgrade paths without restarting the check.")
    add_bullet("Step 6: Official Provenance & Sources", "Collapsible verification drawer citing official manufacturer specs (Adobe, Google, Epic Games, Microsoft) and verification dates.")

    # --- SECTION 7 ---
    add_h1("7. Mobile-First UX, Touch Targets & Quick Actions")
    add_bullet("Minimum 44px Touch Targets", "All interactive buttons, chips, and toggles enforce touch-target sizing with generous tap padding.")
    add_bullet("Sticky Mobile Action Dock", "Fixed floating action bar positioned in the thumb-reachable zone above bottom navigation on mobile viewports.")
    add_bullet("Phone Transfer (QR Code Modal)", "Instant QR code generation allowing users to transfer their exact desktop hardware & workload configuration to their smartphone.")
    add_bullet("Shareable Diagnostic Link", "Share modal generates permanent encoded URLs and markdown summaries for forums, Discord, and technical support.")

    # --- SECTION 8 ---
    add_h1("8. Verification & Test Invariants")
    add_body("The implementation adheres to strict mathematical correctness invariants validated across 30 test suites and 150 automated vitest scenarios:")
    add_bullet("Monotonicity Invariant", "Upgrading a hardware component (e.g. 16GB -> 32GB RAM) is mathematically guaranteed to never decrease a performance score.")
    add_bullet("Concurrency Invariant", "Adding concurrent workloads strictly increases memory pressure ratio and thread contention penalties.")
    add_bullet("Unified vs Dedicated Invariant", "Apple Silicon unified memory correctly models shared RAM/VRAM pools, while discrete GPUs correctly isolate PCIe VRAM from motherboard RAM.")

    doc.save(output_path)
    print("Successfully generated updated design document at:", output_path)

if __name__ == "__main__":
    out_file = "/Users/kerenlint/MyProjects/computebestspecs/ComputeBestSpecs_UI_Design_System_Spec_Updated.docx"
    create_styled_document(out_file)
